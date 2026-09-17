from openpyxl import Workbook, load_workbook
import pytest

from rule_excel import HEADERS, RuleExcelError, export_rules_to_excel, parse_rules_excel
from rule_models import CustomRule


def make_rule(operator="range", **kwargs):
    defaults = dict(column="M", gender="共用", operator=operator, level="第一級", minimum=1, maximum=10, header="BMI", note="醫師確認")
    defaults.update(kwargs)
    rule = CustomRule(**defaults); rule.validate(); return rule


def test_export_format_values_and_validation(tmp_path):
    path = tmp_path / "rules.xlsx"
    rules = [make_rule(), make_rule("exact", column="O", minimum=None, maximum=None, value="++")]
    export_rules_to_excel(path, rules)
    wb = load_workbook(path); ws = wb["規則"]
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:K3"
    assert ws["F2"].value == 1 and ws["G2"].value == 10
    assert ws["H3"].value == "++"
    assert len(ws.data_validations.dataValidation) == 3
    assert "填寫說明" in wb.sheetnames


def write_rows(path, rows, headers=HEADERS):
    wb = Workbook(); ws = wb.active; ws.title = "規則"; ws.append(headers)
    for row in rows: ws.append(row)
    wb.save(path)


def test_import_normalizes_values_preserves_note_and_ignores_blank(tmp_path):
    path = tmp_path / "rules.xlsx"
    write_rows(path, [[" 1 ", "m", "BMI", "共用", ">=", None, None, " 30 ", "第三級", "新版", None], [None] * 11])
    result = parse_rules_excel(path)
    assert not result.has_errors and len(result.rules) == 1
    assert result.rules[0].value == 30
    assert result.rules[0].column == "M" and result.rules[0].note == "新版"


@pytest.mark.parametrize("row,message", [
    ([1, "A1", "BMI", "共用", ">=", None, None, 27, "第三級"], "欄位格式錯誤"),
    ([1, "M", "BMI", "男性", ">=", None, None, 27, "第三級"], "性別"),
    ([1, "M", "BMI", "共用", ">=", None, None, "ABC", "第三級"], "必須是數字"),
    ([1, "M", "BMI", "共用", "區間", 30, 20, None, "第三級"], "下限不可大於上限"),
    ([1, "M", "BMI", "共用", "完全相符", None, None, None, "第三級"], "不可空白"),
])
def test_import_errors(tmp_path, row, message):
    path = tmp_path / "bad.xlsx"; write_rows(path, [row]); result = parse_rules_excel(path)
    assert result.has_errors and any(message in item for item in result.rows[0].errors)


def test_missing_headers_and_conflict_warnings(tmp_path):
    path = tmp_path / "missing.xlsx"; write_rows(path, [], HEADERS[:-3])
    with pytest.raises(RuleExcelError, match="缺少欄位"): parse_rules_excel(path)
    path = tmp_path / "overlap.xlsx"
    write_rows(path, [[1, "M", "BMI", "共用", ">", None, None, 10, "第一級"], [1, "M", "BMI", "共用", ">", None, None, 20, "第二級"]])
    result = parse_rules_excel(path)
    assert all(row.severity == "warning" for row in result.rows)
    assert any("規則重疊" in row.warnings for row in result.rows)
    assert any("順序重複" in row.warnings for row in result.rows)


def test_old_excel_enabled_column_is_compatible(tmp_path):
    old_headers = ("啟用",) + HEADERS
    path = tmp_path / "old.xlsx"
    write_rows(path, [["是", 1, "M", "BMI", "共用", ">=", None, None, 27, "第三級"], ["否", 2, "N", "腰圍", "男", ">=", None, None, 90, "第三級"]], old_headers)
    result = parse_rules_excel(path)
    assert [rule.column for rule in result.rules] == ["M"]
    assert result.ignored_legacy_disabled == 1
