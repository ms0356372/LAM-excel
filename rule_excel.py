"""自訂規則 Excel 交換格式；不依賴 Tkinter 或 Excel COM。"""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rule_models import LABEL_TO_OPERATOR, OPERATOR_LABELS, CustomRule
from rule_view import collect_rule_statuses

RULE_SHEET = "規則"
HEADERS = ("順序", "欄位", "表頭", "性別", "判斷方式", "下限", "上限", "判斷值", "分級", "備註", "檢查結果")
REQUIRED_HEADERS = HEADERS[:9]
class RuleExcelError(ValueError): pass

@dataclass
class ImportRow:
    excel_row: int
    order: int | None = None
    rule: CustomRule | None = None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    @property
    def severity(self): return "error" if self.errors else "warning" if self.warnings else "normal"

@dataclass
class ImportResult:
    rows: list[ImportRow]
    global_warnings: list[str] = field(default_factory=list)
    ignored_legacy_disabled: int = 0
    @property
    def rules(self): return [r.rule for r in sorted(self.rows, key=lambda x: x.order or 0) if r.rule and not r.errors]
    @property
    def has_errors(self): return any(r.errors for r in self.rows)

def _text(value): return "" if value is None else str(value).strip()
def _legacy_enabled(value):
    if isinstance(value, bool): return value
    value = _text(value).casefold()
    if value in {"是", "true", "1"}: return True
    if value in {"否", "false", "0"}: return False
    raise ValueError("舊版啟用欄只允許是／否、TRUE／FALSE 或 1／0。")
def _number(value, label):
    try:
        if value is None or isinstance(value, bool) or not _text(value): raise ValueError
        return float(_text(value))
    except (TypeError, ValueError) as exc: raise ValueError(f"{label}必須是數字。") from exc

def validate_rule_excel_headers(values):
    mapping = {_text(v): i for i, v in enumerate(values) if _text(v)}
    missing = [h for h in REQUIRED_HEADERS if h not in mapping]
    if missing: raise RuleExcelError("Excel 規則檔格式不正確，缺少欄位：\n" + "\n".join(missing))
    return mapping

def parse_excel_rule(values, mapping, excel_row):
    get = lambda name: values[mapping[name]] if name in mapping and mapping[name] < len(values) else None
    row = ImportRow(excel_row)
    try:
        raw_order = get("順序"); row.order = int(raw_order)
        if row.order <= 0 or float(raw_order) != row.order: raise ValueError
    except (TypeError, ValueError): row.errors.append("順序必須為正整數。")
    column, gender, label, level = _text(get("欄位")).upper(), _text(get("性別")), _text(get("判斷方式")), _text(get("分級"))
    operator = LABEL_TO_OPERATOR.get(label, label); minimum = maximum = value = None
    if label == "區間":
        try: minimum = _number(get("下限"), "區間下限")
        except ValueError as exc: row.errors.append(str(exc))
        try: maximum = _number(get("上限"), "區間上限")
        except ValueError as exc: row.errors.append(str(exc))
        if minimum is not None and maximum is not None and minimum > maximum: row.errors.append("區間下限不可大於上限。")
        if _text(get("判斷值")): row.errors.append("區間規則的判斷值應為空白。")
    elif label in {">", ">=", "<", "<="}:
        try: value = _number(get("判斷值"), "比較數值")
        except ValueError as exc: row.errors.append(str(exc))
        if _text(get("下限")) or _text(get("上限")): row.errors.append("數值比較規則的下限與上限應為空白。")
    elif label == "完全相符":
        value = _text(get("判斷值"))
        if not value: row.errors.append("完全相符的內容不可空白。")
        if _text(get("下限")) or _text(get("上限")): row.errors.append("完全相符規則的下限與上限應為空白。")
    rule = CustomRule(column, gender, operator, level, value, minimum, maximum, _text(get("表頭")), _text(get("備註")))
    try: rule.validate()
    except ValueError as exc:
        if str(exc) not in row.errors: row.errors.append(str(exc))
    if not rule.header: row.warnings.append("表頭空白")
    row.rule = rule
    return row

def parse_rules_excel(path):
    try: wb = load_workbook(path, read_only=True, data_only=True)
    except PermissionError as exc: raise RuleExcelError("無法讀取規則 Excel，請確認檔案是否正在使用中。") from exc
    except Exception as exc: raise RuleExcelError(f"無法讀取規則 Excel，檔案可能已損毀：{exc}") from exc
    try:
        if RULE_SHEET not in wb.sheetnames: raise RuleExcelError("Excel 規則檔找不到「規則」工作表。")
        ws = wb[RULE_SHEET]
        mapping = validate_rule_excel_headers([c.value for c in next(ws.iter_rows(min_row=1, max_row=1))])
        rows = []; ignored = 0
        for number, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            values = list(cells)
            if all(not _text(v) for v in values): continue
            if "啟用" in mapping:
                try: legacy_enabled = _legacy_enabled(values[mapping["啟用"]] if mapping["啟用"] < len(values) else None)
                except ValueError as exc:
                    rows.append(ImportRow(number, errors=[str(exc)])); continue
                if not legacy_enabled:
                    ignored += 1; continue
            rows.append(parse_excel_rule(values, mapping, number))
    finally: wb.close()
    if not rows: raise RuleExcelError("規則工作表中沒有可匯入的規則。")
    result = ImportResult(rows, ignored_legacy_disabled=ignored); _order_warnings(result)
    valid = [r for r in rows if r.rule and not r.errors]
    for row, status in zip(valid, collect_rule_statuses([r.rule for r in valid])):
        for message in status.messages:
            if message not in row.warnings: row.warnings.append(message)
    return result

def _order_warnings(result):
    valid = [r for r in result.rows if r.order is not None]
    counts = {o: sum(r.order == o for r in valid) for o in {r.order for r in valid}}
    for row in valid:
        if counts[row.order] > 1: row.warnings.append("順序重複，匯入後將重新編號")
    if counts and sorted(counts) != list(range(1, len(counts) + 1)):
        result.global_warnings.append("規則順序不連續，匯入後將重新編號。")
        for row in valid: row.warnings.append("規則順序不連續，匯入後將重新編號")

def format_rule_for_excel(rule, order, status, headers=None):
    return [order, rule.column, rule.header or (headers or {}).get(rule.column, ""), rule.gender,
            OPERATOR_LABELS[rule.operator], rule.minimum if rule.operator == "range" else None,
            rule.maximum if rule.operator == "range" else None, rule.value if rule.operator != "range" else None,
            rule.level, rule.note, status]

def export_rules_to_excel(path, rules, headers=None):
    statuses = collect_rule_statuses(rules); wb = Workbook(); ws = wb.active; ws.title = RULE_SHEET; ws.append(HEADERS)
    for i, (rule, status) in enumerate(zip(rules, statuses), 1): ws.append(format_rule_for_excel(rule, i, "正常" if not status.messages else "；".join(status.messages), headers))
    _format_sheet(ws); _instructions(wb)
    try: wb.save(path)
    except PermissionError as exc: raise RuleExcelError("無法儲存規則 Excel。\n請確認檔案是否正在 Excel 中開啟，或選擇其他輸出位置。") from exc

def create_rule_excel_template(path): export_rules_to_excel(path, [])
def _format_sheet(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]: c.font = Font(bold=True); c.fill = fill; c.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:K{max(ws.max_row, 1)}"
    for i, width in enumerate((9, 9, 18, 10, 14, 12, 12, 18, 12, 24, 30), 1): ws.column_dimensions[get_column_letter(i)].width = width
    for col, options in (("D", '"共用,男,女"'), ("E", '"區間,>,>=,<,<=,完全相符"'), ("I", '"第一級,第二級,第三級,第四級"')):
        dv = DataValidation(type="list", formula1=options, allow_blank=False); dv.error = "請從下拉清單選擇有效值。"; dv.showErrorMessage = True
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}10000")
def _instructions(wb):
    ws = wb.create_sheet("填寫說明")
    lines = ["規則檔版本：2", "", "【順序】決定 first match wins 的判斷順序，數字越小越先判斷。", "【欄位】填 Excel 欄位字母，例如 M、N、AA。", "【表頭】建議填第一列表頭名稱，方便人工確認。", "【性別】共用、男、女", "【判斷方式】區間、>、>=、<、<=、完全相符", "【區間】下限 18.5、上限 23.9 代表 18.5 <= value <= 23.9。", "【大於】> 10.6，10.6 不符合。", "【大於等於】>= 10.6，10.6 符合。", "【小於】< 25.8，25.8 不符合。", "【小於等於】<= 25.8，25.8 符合。", "【完全相符】例如 ++、代謝症候群，必須完全一致。", "【分級】第一級、第二級、第三級、第四級", "【舊版相容】舊檔若有啟用欄，是／TRUE／1 會匯入；否／FALSE／0 會忽略。", "【注意事項】", "• 存在的規則會直接生效；不再使用的規則請在程式中刪除。", "• 空白列不會匯入。", "• 數值規則必須使用數字。", "• 規則順序會影響 first match wins。", "• 男／女專用規則優先於共用規則。", "• 規則重疊時會顯示警告。", "• Excel 匯入前會先預覽，不會直接覆蓋規則。"]
    for line in lines: ws.append([line])
    ws.column_dimensions["A"].width = 90; ws["A1"].font = Font(bold=True)
